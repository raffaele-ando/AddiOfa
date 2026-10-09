"""Genera AddiOFA_conti.xlsx: ipotesi, conto del 2027, mesi, tre anni, pareggio.

Ogni numero in blu è un'ipotesi da cambiare nel foglio Ipotesi; tutto il resto sono formule.
Lanciare con: python3 docs/business/genera_conti.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

QUI = Path(__file__).resolve().parent
USCITA = QUI / "AddiOFA_conti.xlsx"

FONT = "Arial"
BLU = Font(name=FONT, color="0000FF")
NERO = Font(name=FONT)
VERDE = Font(name=FONT, color="008000")
GRASSETTO = Font(name=FONT, bold=True)
TITOLO = Font(name=FONT, bold=True, size=14)
GIALLO = PatternFill("solid", fgColor="FFFF00")
GRIGIO = PatternFill("solid", fgColor="EEEEEE")
SOTTILE = Side(style="thin", color="BBBBBB")
BORDO_SOPRA = Border(top=Side(style="thin", color="000000"))

EURO = '#,##0 "€";(#,##0 "€");"-"'
EURO2 = '#,##0.00 "€";(#,##0.00 "€");"-"'
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'

SCENARI = ["Prudente", "Base", "Buono"]
COL = {"Prudente": "C", "Base": "D", "Buono": "E"}

# ---------------------------------------------------------------- ipotesi
# (chiave, etichetta, formato, valori per scenario, fonte/nota, chiave?)
IPOTESI = [
    ("sez", "Mercato del recupero (studenti con OFA di inglese)"),
    ("matricole", "Immatricolati l'anno ai corsi in italiano", NUM, [7000, 7000, 7000],
     "USTAT: 7.820 immatricolati 2024/25, meno circa 800 nei corsi in inglese (stima)."),
    ("q_ofa", "Quota con OFA di inglese", PCT, [0.21, 0.22, 0.23],
     "Graduatorie 2022-2026: 21-25 % degli ammessi con il dato; sondaggio Gestionale 23,5 %."),
    ("q_cert", "Quota che lo toglie con un certificato B1 già posseduto", PCT, [0.30, 0.22, 0.15],
     "Ipotesi: non esiste un dato pubblico. Da misurare con un sondaggio."),
    ("arretrati", "Studenti degli anni precedenti ancora con l'OFA", NUM, [100, 200, 300],
     "Ipotesi."),
    ("q_prova", "Quota del mercato che finisce il diagnostico gratuito", PCT, [0.25, 0.40, 0.55],
     "Ipotesi chiave: dipende da accordo PoliNetwork, Google e passaparola.", True),
    ("conv_rec", "Conversione a pagamento di chi ha finito il diagnostico", PCT, [0.05, 0.09, 0.13],
     "RevenueCat 2026: mediana 2,1 %, quartile alto 4,5 % sui download; qui il bisogno è acuto.", True),
    ("sez", "Mercato della prevenzione (candidati al test d'ingresso)"),
    ("candidati", "Candidati l'anno al test d'ingresso (TENG incluso)", NUM, [14000, 14000, 14000],
     "Graduatorie 2026: 14.162 candidati unici."),
    ("q_prev", "Quota dei candidati che prova l'app", PCT, [0.01, 0.025, 0.04],
     "Ipotesi: solo contenuti gratuiti (YouTube, Google), primo anno.", True),
    ("conv_prev", "Conversione a pagamento dei candidati che provano", PCT, [0.03, 0.05, 0.08],
     "Ipotesi: bisogno meno urgente del recupero."),
    ("sez", "Prezzi e mix"),
    ("prezzo", "Prezzo del Pass", EURO2, [14.99, 14.99, 14.99], "Proposta del documento."),
    ("prezzo_lancio", "Prezzo di lancio", EURO2, [9.99, 9.99, 9.99], "Fino al 31/01/2027, scadenza vera."),
    ("q_lancio", "Quota delle vendite dell'anno a prezzo di lancio", PCT, [0.30, 0.20, 0.15], "Ipotesi."),
    ("q_invitati", "Quota delle vendite fatte da invitati (sconto 20 %)", PCT, [0.15, 0.20, 0.25],
     "Andrew Chen: il passaparola aggiunge il 25-43 % di utenti."),
    ("sconto_inv", "Sconto per gli invitati", PCT, [0.20, 0.20, 0.20], "Proposta del documento."),
    ("q_garanzia", "Quota dei Pass venduti con garanzia (+5 €)", PCT, [0.10, 0.15, 0.20],
     "Ipotesi: la garanzia parte a giugno 2027."),
    ("premio_gar", "Sovrapprezzo della garanzia", EURO2, [5, 5, 5], "19,99 contro 14,99."),
    ("rimb_gar", "Quota dei Pass con garanzia rimborsati", PCT, [0.10, 0.07, 0.05],
     "Ipotesi: si misura nella prima estate."),
    ("q_amb", "Quota delle vendite portate da ambassador", PCT, [0.10, 0.15, 0.20], "Ipotesi."),
    ("comm_amb", "Commissione degli ambassador", PCT, [0.20, 0.20, 0.20], "Proposta del documento."),
    ("sessioni", "Corsi live l'anno", NUM, [0, 1, 2], "Solo se la lista d'attesa supera 15 iscritti."),
    ("iscritti", "Iscritti per corso live", NUM, [10, 15, 25], "Ipotesi."),
    ("prezzo_corso", "Prezzo del corso live", EURO2, [25, 25, 29], "Proposta del documento."),
    ("costo_docente", "Costo del docente per corso (0 se lo tieni tu)", EURO, [0, 80, 80], "Ipotesi: 2 ore più preparazione."),
    ("sez", "Pagamenti e perdite"),
    ("stripe_pct", "Commissione Stripe in percentuale", PCT, [0.015, 0.015, 0.015], "stripe.com/it/pricing: carte UE standard."),
    ("stripe_fisso", "Commissione Stripe fissa per pagamento", EURO2, [0.25, 0.25, 0.25], "stripe.com/it/pricing."),
    ("q_recessi", "Recessi e contestazioni sui ricavi", PCT, [0.04, 0.03, 0.02], "Benchmark: 2-5 % (ricerca)."),
    ("q_polinetwork", "Quota dei ricavi alla community (accordo PoliNetwork)", PCT, [0.05, 0.05, 0.05],
     "Proposta da negoziare; 0 se non c'è accordo."),
    ("sez", "Costi fissi annui"),
    ("c_commercialista", "Commercialista o servizio online per forfettari", EURO, [350, 350, 350],
     "Listini 2026: 199-499 € l'anno (TaxMan, ForfettApp, FlexTax, Fiscozen)."),
    ("c_camera", "Diritto camerale (solo se iscritto al Registro imprese)", EURO, [0, 0, 0],
     "Circa 53 € l'anno se commerciante; 0 se professionista."),
    ("c_pec", "PEC e firma digitale", EURO, [40, 40, 40], "Ipotesi."),
    ("c_dominio", "Dominio .it", EURO, [15, 15, 15], "Aruba: 11,99 € + IVA."),
    ("c_cloud", "Cloudflare (Workers a pagamento quando servono)", EURO, [0, 60, 60], "5 $ al mese; gratis fino a 100.000 richieste al giorno."),
    ("c_legale", "Termini, privacy e cookie", EURO, [60, 60, 60], "Generatore a pagamento più controllo; ipotesi."),
    ("c_stampa", "Volantini e stampa", EURO, [50, 80, 100], "PressUP: 1.000 A6 da 49 € + IVA."),
    ("c_altro", "Strumenti vari ed email", EURO, [0, 30, 60], "Resend gratis fino a 3.000 email al mese."),
    ("sez", "Fisco e contributi (da confermare con un commercialista)"),
    ("c_impresa", "Costi fissi in più se è un'impresa (camera di commercio, apertura, commercialista più caro)", EURO, [300, 300, 300],
     "Diritto camerale 53 €, apertura ditta circa 150 € il primo anno, Fiscozen 599 € contro 499 €."),
    ("coeff", "Coefficiente di redditività del regime forfettario", PCT, [0.67, 0.67, 0.67],
     "ATECO 2025 58.29.00 (edizione di software) o 62.10.00: 67 %. Corsi 85.59.20: 78 %.", True),
    ("imposta", "Imposta sostitutiva (primi 5 anni)", PCT, [0.05, 0.05, 0.05], "L. 190/2014, regime forfettario startup."),
    ("aliq_gs", "Aliquota INPS Gestione Separata", PCT, [0.2607, 0.2607, 0.2607], "Circolare INPS 8/2026."),
    ("min_com", "Contributo minimo annuo INPS Commercianti (prima della riduzione)", EURO, [4611.64, 4611.64, 4611.64],
     "Circolare INPS 14/2026: dovuto anche con incassi zero."),
    ("aliq_com", "Aliquota INPS Commercianti oltre il minimale", PCT, [0.2448, 0.2448, 0.2448], "Circolare INPS 14/2026."),
    ("redd_min_com", "Reddito minimale INPS Commercianti", EURO, [18808, 18808, 18808], "Circolare INPS 14/2026."),
    ("rid_forf", "Riduzione dei contributi Commercianti per i forfettari", PCT, [0.35, 0.35, 0.35], "Su domanda all'INPS (L. 190/2014)."),
    ("sez", "Obiettivo personale"),
    ("obiettivo", "Netto annuo per vivere a Milano", EURO, [16400, 16400, 16400],
     "Stima della chat: affitto, cibo, svago, vestiti, retta."),
]

# ---------------------------------------------------------------- crescita (foglio Tre anni)
CRESCITA = [
    ("cr_prova", "Aumento della quota che prova l'app, anno su anno", PCT, [0.10, 0.25, 0.40]),
    ("cr_prev", "Moltiplicatore della quota di candidati che prova, anno su anno", '0.0"x"', [1.5, 2.0, 2.5]),
    ("nuovo_ateneo", "Ricavi lordi da un secondo mercato dal 3° anno (es. Bicocca o TOLC)", EURO, [0, 2000, 5000]),
]

wb = Workbook()

# ---------------------------------------------------------------- Leggimi
ws = wb.active
ws.title = "Leggimi"
righe = [
    ("AddiOFA: conti del primo anno e dei tre anni", TITOLO),
    ("", NERO),
    ("Come si usa", GRASSETTO),
    ("1. Le ipotesi sono tutte nel foglio Ipotesi, in tre colonne: Prudente, Base, Buono.", NERO),
    ("2. Numeri in blu = ipotesi da cambiare. Numeri in nero = formule: non toccarli.", NERO),
    ("3. Le celle gialle sono le ipotesi che pesano di più: misurale con la beta e aggiornale.", NERO),
    ("4. Nel foglio Ipotesi, cella C3, scegli lo scenario mostrato nei fogli Mesi e Tre anni.", NERO),
    ("5. Il risultato è calcolato in due modi: Gestione Separata (professionista) e Commercianti (impresa). Quale vale lo decide il commercialista.", NERO),
    ("", NERO),
    ("Cosa contiene", GRASSETTO),
    ("Ipotesi: tutte le leve, con la fonte o la nota accanto.", NERO),
    ("Anno 2027: conto economico del primo anno di vendita (gennaio-dicembre 2027) nei tre scenari.", NERO),
    ("Mesi: come si distribuiscono incassi e costi nei 12 mesi del 2027, scenario scelto.", NERO),
    ("Tre anni: 2027-2029 con le ipotesi di crescita.", NERO),
    ("Pareggio: quanti Pass servono per coprire i costi fissi, con le due gestioni INPS.", NERO),
    ("", NERO),
    ("Avvertenze", GRASSETTO),
    ("Le ipotesi di conversione non sono dati misurati: lo diventano solo con la beta.", NERO),
    ("Il calcolo fiscale è semplificato (forfettario, cassa = competenza, senza IVA perché il forfettario non la applica).", NERO),
    ("Va confermato con un commercialista prima del primo incasso.", NERO),
]
for i, (testo, font) in enumerate(righe, start=1):
    c = ws.cell(row=i, column=1, value=testo)
    c.font = font
ws.column_dimensions["A"].width = 110

# ---------------------------------------------------------------- Ipotesi
ip = wb.create_sheet("Ipotesi")
ip["A1"] = "Ipotesi"
ip["A1"].font = TITOLO
ip["A3"] = "Scenario mostrato nei fogli Mesi e Tre anni"
ip["A3"].font = GRASSETTO
ip["C3"] = "Base"
ip["C3"].font = BLU
ip["C3"].fill = GIALLO
dv = DataValidation(type="list", formula1='"Prudente,Base,Buono"', allow_blank=False)
ip.add_data_validation(dv)
dv.add("C3")
ip["D3"] = '=IF(C3="Prudente",1,IF(C3="Buono",3,2))'
ip["D3"].font = NERO
ip["E3"] = "← numero dello scenario (1-3), usato dalle formule"
ip["E3"].font = Font(name=FONT, italic=True, color="666666")

intest = 5
for j, t in enumerate(["Ipotesi", "", "Prudente", "Base", "Buono", "Fonte o nota"], start=1):
    c = ip.cell(row=intest, column=j, value=t)
    c.font = GRASSETTO
    c.fill = GRIGIO

RIGA = {}  # chiave -> riga nel foglio Ipotesi
r = intest + 1
for voce in IPOTESI:
    if voce[0] == "sez":
        r += 1
        ip.cell(row=r, column=1, value=voce[1]).font = GRASSETTO
        r += 1
        continue
    chiave, etichetta, fmt, valori, nota = voce[:5]
    chiave_forte = len(voce) > 5 and voce[5]
    ip.cell(row=r, column=1, value=etichetta).font = NERO
    for k, sc in enumerate(SCENARI):
        c = ip[f"{COL[sc]}{r}"]
        c.value = valori[k]
        c.font = BLU
        if fmt:
            c.number_format = fmt
        if chiave_forte:
            c.fill = GIALLO
    ip.cell(row=r, column=6, value=nota).font = Font(name=FONT, color="444444")
    RIGA[chiave] = r
    r += 1

r += 1
ip.cell(row=r, column=1, value="Crescita (foglio Tre anni)").font = GRASSETTO
r += 1
for chiave, etichetta, fmt, valori in CRESCITA:
    ip.cell(row=r, column=1, value=etichetta).font = NERO
    for k, sc in enumerate(SCENARI):
        c = ip[f"{COL[sc]}{r}"]
        c.value = valori[k]
        c.font = BLU
        c.number_format = fmt
    ip.cell(row=r, column=6, value="Ipotesi.").font = Font(name=FONT, color="444444")
    RIGA[chiave] = r
    r += 1


ip.column_dimensions["A"].width = 62
ip.column_dimensions["B"].width = 2
for c in "CDE":
    ip.column_dimensions[c].width = 13
ip.column_dimensions["F"].width = 95
ip.freeze_panes = "C6"


def I(chiave, col):
    """Riferimento a un'ipotesi per una colonna di scenario."""
    return f"Ipotesi!${col}${RIGA[chiave]}"


def Isel(chiave):
    """Riferimento a un'ipotesi nello scenario scelto in Ipotesi!C3."""
    return f"INDEX(Ipotesi!$C${RIGA[chiave]}:$E${RIGA[chiave]},Ipotesi!$D$3)"


# ---------------------------------------------------------------- Anno 2027
an = wb.create_sheet("Anno 2027")
an["A1"] = "Conto del primo anno di vendita: gennaio-dicembre 2027"
an["A1"].font = TITOLO
an["A2"] = "Tutte le celle sono formule sulle ipotesi. Importi in euro."
an["A2"].font = Font(name=FONT, italic=True, color="666666")
for j, t in enumerate(["Voce", "", "Prudente", "Base", "Buono"], start=1):
    c = an.cell(row=4, column=j, value=t)
    c.font = GRASSETTO
    c.fill = GRIGIO

# (chiave, etichetta, formato, funzione col -> formula, grassetto?)
def f_mercato(c):
    return f"={I('matricole', c)}*{I('q_ofa', c)}*(1-{I('q_cert', c)})+{I('arretrati', c)}"


VOCI = [
    ("sez", "Clienti"),
    ("mercato", "Mercato del recupero (persone)", NUM, f_mercato, False),
    ("provano", "Finiscono il diagnostico", NUM, lambda c: f"={{mercato}}*{I('q_prova', c)}", False),
    ("pag_rec", "Comprano il Pass (recupero)", NUM, lambda c: f"={{provano}}*{I('conv_rec', c)}", False),
    ("prov_prev", "Candidati che provano l'app (prevenzione)", NUM, lambda c: f"={I('candidati', c)}*{I('q_prev', c)}", False),
    ("pag_prev", "Comprano il Pass (prevenzione)", NUM, lambda c: f"={{prov_prev}}*{I('conv_prev', c)}", False),
    ("paganti", "Pass venduti in totale", NUM, lambda c: "={pag_rec}+{pag_prev}", True),
    ("quota_merc", "Quota del mercato del recupero che compra", PCT, lambda c: "=IF({mercato}=0,0,{pag_rec}/{mercato})", False),
    ("sez", "Ricavi"),
    ("prezzo_medio", "Prezzo medio incassato per Pass, senza garanzia", EURO2,
     lambda c: f"=({I('q_lancio', c)}*{I('prezzo_lancio', c)}+(1-{I('q_lancio', c)})*{I('prezzo', c)})*(1-{I('q_invitati', c)}*{I('sconto_inv', c)})", False),
    ("ric_pass", "Ricavi dai Pass", EURO, lambda c: "={paganti}*{prezzo_medio}", False),
    ("ric_gar", "Ricavi dal sovrapprezzo della garanzia", EURO, lambda c: f"={{paganti}}*{I('q_garanzia', c)}*{I('premio_gar', c)}", False),
    ("ric_corso", "Ricavi dai corsi live", EURO, lambda c: f"={I('sessioni', c)}*{I('iscritti', c)}*{I('prezzo_corso', c)}", False),
    ("ricavi", "Ricavi lordi", EURO, lambda c: "={ric_pass}+{ric_gar}+{ric_corso}", True),
    ("sez", "Costi variabili"),
    ("c_stripe", "Commissioni Stripe", EURO,
     lambda c: f"={{ricavi}}*{I('stripe_pct', c)}+({{paganti}}+{I('sessioni', c)}*{I('iscritti', c)})*{I('stripe_fisso', c)}", False),
    ("c_recessi", "Recessi e contestazioni", EURO, lambda c: f"={{ricavi}}*{I('q_recessi', c)}", False),
    ("c_rimborsi", "Rimborsi della garanzia", EURO,
     lambda c: f"={{paganti}}*{I('q_garanzia', c)}*{I('rimb_gar', c)}*({{prezzo_medio}}+{I('premio_gar', c)})", False),
    ("c_amb", "Commissioni agli ambassador", EURO, lambda c: f"={{ric_pass}}*{I('q_amb', c)}*{I('comm_amb', c)}", False),
    ("c_pn", "Quota alla community", EURO, lambda c: f"={{ricavi}}*{I('q_polinetwork', c)}", False),
    ("c_docente", "Docenti dei corsi live", EURO, lambda c: f"={I('sessioni', c)}*{I('costo_docente', c)}", False),
    ("c_var", "Totale costi variabili", EURO, lambda c: "={c_stripe}+{c_recessi}+{c_rimborsi}+{c_amb}+{c_pn}+{c_docente}", True),
    ("sez", "Costi fissi"),
    ("c_fissi", "Totale costi fissi", EURO,
     lambda c: "=" + "+".join(I(k, c) for k in ["c_commercialista", "c_camera", "c_pec", "c_dominio", "c_cloud", "c_legale", "c_stampa", "c_altro"]), True),
    ("sez", "Risultato"),
    ("margine", "Margine prima di contributi e imposta", EURO, lambda c: "={ricavi}-{c_var}-{c_fissi}", True),
    ("ric_fisc", "Ricavi fiscali (incassi meno rimborsi e recessi)", EURO, lambda c: "={ricavi}-{c_recessi}-{c_rimborsi}", False),
    ("reddito", "Reddito forfettario (ricavi fiscali × coefficiente)", EURO, lambda c: f"={{ric_fisc}}*{I('coeff', c)}", False),
    ("inps_gs", "Contributi INPS se Gestione Separata (professionista)", EURO,
     lambda c: f"={{reddito}}*{I('aliq_gs', c)}", False),
    ("imp_gs", "Imposta sostitutiva se Gestione Separata", EURO,
     lambda c: f"=MAX(0,{{reddito}}-{{inps_gs}})*{I('imposta', c)}", False),
    ("netto", "Utile netto in tasca se Gestione Separata", EURO, lambda c: "={margine}-{inps_gs}-{imp_gs}", True),
    ("inps_com", "Contributi INPS se Commercianti (impresa)", EURO,
     lambda c: (f"=(MAX({I('min_com', c)},{I('min_com', c)}+MAX(0,{{reddito}}-{I('redd_min_com', c)})*{I('aliq_com', c)}))*(1-{I('rid_forf', c)})"), False),
    ("imp_com", "Imposta sostitutiva se Commercianti", EURO,
     lambda c: f"=MAX(0,{{reddito}}-{{inps_com}})*{I('imposta', c)}", False),
    ("netto_com", "Utile netto in tasca se Commercianti", EURO,
     lambda c: f"={{margine}}-{I('c_impresa', c)}-{{inps_com}}-{{imp_com}}", True),
    ("netto_mese", "Utile netto medio al mese se Gestione Separata", EURO, lambda c: "={netto}/12", False),
    ("copertura", "Quota dell'obiettivo personale coperta (Gestione Separata)", PCT, lambda c: f"={{netto}}/{I('obiettivo', c)}", True),
]

POS = {}
r = 5
righe_formule = []
for voce in VOCI:
    if voce[0] == "sez":
        r += 1
        an.cell(row=r, column=1, value=voce[1]).font = GRASSETTO
        r += 1
        continue
    chiave, etichetta, fmt, fn, forte = voce
    POS[chiave] = r
    righe_formule.append((r, chiave, etichetta, fmt, fn, forte))
    r += 1

for r, chiave, etichetta, fmt, fn, forte in righe_formule:
    an.cell(row=r, column=1, value=etichetta).font = GRASSETTO if forte else NERO
    for sc in SCENARI:
        col = COL[sc]
        formula = fn(col).format(**{k: f"{col}{v}" for k, v in POS.items()})
        c = an[f"{col}{r}"]
        c.value = formula
        c.number_format = fmt
        c.font = GRASSETTO if forte else NERO
        if forte:
            c.border = BORDO_SOPRA
an.column_dimensions["A"].width = 55
an.column_dimensions["B"].width = 2
for c in "CDE":
    an.column_dimensions[c].width = 14
an.freeze_panes = "C5"


def A(chiave, col):
    return f"'Anno 2027'!${col}${POS[chiave]}"


def Asel(chiave):
    return f"INDEX('Anno 2027'!$C${POS[chiave]}:$E${POS[chiave]},Ipotesi!$D$3)"


# ---------------------------------------------------------------- Mesi
me = wb.create_sheet("Mesi")
me["A1"] = "Il 2027 mese per mese, scenario scelto in Ipotesi!C3"
me["A1"].font = TITOLO
me["A2"] = "I pesi (in blu) dicono in che mese cadono le vendite: devono sommare a 100 %."
me["A2"].font = Font(name=FONT, italic=True, color="666666")
mesi = ["gen", "feb", "mar", "apr", "mag", "giu", "lug", "ago", "set", "ott", "nov", "dic"]
peso_rec = [0.10, 0.14, 0.06, 0.02, 0.02, 0.05, 0.14, 0.12, 0.22, 0.08, 0.04, 0.01]
peso_prev = [0.10, 0.20, 0.25, 0.25, 0.20, 0, 0, 0, 0, 0, 0, 0]
peso_corso = [0.5, 0, 0, 0, 0, 0, 0.5, 0, 0, 0, 0, 0]
me["A4"] = "Voce"
me["A4"].font = GRASSETTO
me["A4"].fill = GRIGIO
for j, m in enumerate(mesi):
    c = me.cell(row=4, column=3 + j, value=m)
    c.font = GRASSETTO
    c.fill = GRIGIO
    c.alignment = Alignment(horizontal="center")
tot_col = get_column_letter(3 + 12)
me[f"{tot_col}4"] = "Totale"
me[f"{tot_col}4"].font = GRASSETTO
me[f"{tot_col}4"].fill = GRIGIO

righe_mesi = [
    (5, "Peso delle vendite di recupero", peso_rec, PCT),
    (6, "Peso delle vendite di prevenzione", peso_prev, PCT),
    (7, "Peso dei corsi live", peso_corso, PCT),
]
for rr, et, pesi, fmt in righe_mesi:
    me.cell(row=rr, column=1, value=et).font = NERO
    for j, p in enumerate(pesi):
        c = me.cell(row=rr, column=3 + j, value=p)
        c.font = BLU
        c.number_format = fmt
    me[f"{tot_col}{rr}"] = f"=SUM(C{rr}:{get_column_letter(14)}{rr})"
    me[f"{tot_col}{rr}"].number_format = PCT

# ricavi del mese: quota recupero e prevenzione dei Pass (con garanzia distribuita come i Pass)
calcoli = [
    (9, "Pass venduti", NUM,
     lambda L: f"={Asel('pag_rec')}*{L}5+{Asel('pag_prev')}*{L}6"),
    (10, "Ricavi lordi", EURO,
     lambda L: f"=({Asel('ric_pass')}+{Asel('ric_gar')})*IF({Asel('paganti')}=0,0,{L}9/{Asel('paganti')})+{Asel('ric_corso')}*{L}7"),
    (11, "Costi variabili", EURO,
     lambda L: f"=IF({Asel('ricavi')}=0,0,{Asel('c_var')}*{L}10/{Asel('ricavi')})"),
    (12, "Costi fissi (un dodicesimo)", EURO, lambda L: f"={Asel('c_fissi')}/12"),
    (13, "Margine del mese", EURO, lambda L: f"={L}10-{L}11-{L}12"),
    (14, "Margine cumulato", EURO, lambda L: f"=SUM($C$13:{L}13)"),
]
for rr, et, fmt, fn in calcoli:
    me.cell(row=rr, column=1, value=et).font = GRASSETTO if rr in (10, 13) else NERO
    for j in range(12):
        L = get_column_letter(3 + j)
        c = me.cell(row=rr, column=3 + j, value=fn(L))
        c.number_format = fmt
        c.font = NERO
    if rr != 14:
        me[f"{tot_col}{rr}"] = f"=SUM(C{rr}:{get_column_letter(14)}{rr})"
        me[f"{tot_col}{rr}"].number_format = fmt
        me[f"{tot_col}{rr}"].font = GRASSETTO
me["A16"] = "Contributi e imposta si pagano a saldo e acconto (giugno e novembre): qui non sono distribuiti per mese."
me["A16"].font = Font(name=FONT, italic=True, color="666666")
me.column_dimensions["A"].width = 36
me.column_dimensions["B"].width = 2
for j in range(13):
    me.column_dimensions[get_column_letter(3 + j)].width = 10

# ---------------------------------------------------------------- Tre anni
tr = wb.create_sheet("Tre anni")
tr["A1"] = "2027-2029, scenario scelto in Ipotesi!C3"
tr["A1"].font = TITOLO
tr["A2"] = ("Ogni anno cresce la quota del mercato che prova l'app e la quota dei candidati raggiunti; "
            "dal 2029 si può aggiungere un secondo mercato. Prezzo di lancio solo nel 2027.")
tr["A2"].font = Font(name=FONT, italic=True, color="666666")
for j, t in enumerate(["Voce", "", "2027", "2028", "2029"], start=1):
    c = tr.cell(row=4, column=j, value=t)
    c.font = GRASSETTO
    c.fill = GRIGIO

def prezzo_medio_anno(primo):
    lancio = f"{Isel('q_lancio')}" if primo else "0"
    return (f"(({lancio})*{Isel('prezzo_lancio')}+(1-({lancio}))*{Isel('prezzo')})"
            f"*(1-{Isel('q_invitati')}*{Isel('sconto_inv')})")

T = [
    (5, "Quota del mercato del recupero che prova l'app", PCT,
     [f"={Isel('q_prova')}", f"=MIN(0.9,C5*(1+{Isel('cr_prova')}))", f"=MIN(0.9,D5*(1+{Isel('cr_prova')}))"]),
    (6, "Quota dei candidati che prova l'app", PCT,
     [f"={Isel('q_prev')}", f"=MIN(0.5,C6*{Isel('cr_prev')})", f"=MIN(0.5,D6*{Isel('cr_prev')})"]),
    (7, "Pass venduti (recupero)", NUM,
     [f"={Asel('mercato')}*{L}5*{Isel('conv_rec')}" for L in "CDE"]),
    (8, "Pass venduti (prevenzione)", NUM,
     [f"={Isel('candidati')}*{L}6*{Isel('conv_prev')}" for L in "CDE"]),
    (9, "Prezzo medio per Pass", EURO2,
     [f"={prezzo_medio_anno(True)}", f"={prezzo_medio_anno(False)}", f"={prezzo_medio_anno(False)}"]),
    (10, "Ricavi lordi (Pass, garanzia, corsi)", EURO,
     [f"=({L}7+{L}8)*({L}9+{Isel('q_garanzia')}*{Isel('premio_gar')})+{Asel('ric_corso')}" for L in "CDE"]),
    (11, "Ricavi da un secondo mercato", EURO,
     ["=0", "=0", f"={Isel('nuovo_ateneo')}"]),
    (12, "Ricavi lordi totali", EURO, [f"={L}10+{L}11" for L in "CDE"]),
    (13, "Costi variabili (stessa quota sui ricavi del 2027)", EURO,
     [f"=IF({Asel('ricavi')}=0,0,{L}12*{Asel('c_var')}/{Asel('ricavi')})" for L in "CDE"]),
    (14, "Costi fissi", EURO, [f"={Asel('c_fissi')}" for _ in "CDE"]),
    (15, "Reddito forfettario", EURO,
     [f"=({L}12*(1-{Isel('q_recessi')}))*{Isel('coeff')}" for L in "CDE"]),
    (16, "Contributi INPS (Gestione Separata)", EURO, [f"={L}15*{Isel('aliq_gs')}" for L in "CDE"]),
    (17, "Imposta sostitutiva", EURO, [f"=MAX(0,{L}15-{L}16)*{Isel('imposta')}" for L in "CDE"]),
    (18, "Utile netto in tasca (Gestione Separata)", EURO, [f"={L}12-{L}13-{L}14-{L}16-{L}17" for L in "CDE"]),
    (19, "Quota dell'obiettivo personale coperta", PCT, [f"={L}18/{Isel('obiettivo')}" for L in "CDE"]),
    (20, "Utile netto se Commercianti (impresa)", EURO,
     [(f"={L}12-{L}13-{L}14-{Isel('c_impresa')}-(MAX({Isel('min_com')},{Isel('min_com')}+MAX(0,{L}15-{Isel('redd_min_com')})*{Isel('aliq_com')}))*(1-{Isel('rid_forf')})"
       f"-MAX(0,{L}15-(MAX({Isel('min_com')},{Isel('min_com')}+MAX(0,{L}15-{Isel('redd_min_com')})*{Isel('aliq_com')}))*(1-{Isel('rid_forf')}))*{Isel('imposta')}") for L in "CDE"]),
]
for rr, et, fmt, formule in T:
    forte = rr in (12, 18, 19, 20)
    tr.cell(row=rr, column=1, value=et).font = GRASSETTO if forte else NERO
    for j, f in enumerate(formule):
        c = tr.cell(row=rr, column=3 + j, value=f)
        c.number_format = fmt
        c.font = GRASSETTO if forte else NERO
tr.column_dimensions["A"].width = 52
tr.column_dimensions["B"].width = 2
for c in "CDE":
    tr.column_dimensions[c].width = 14

# ---------------------------------------------------------------- Pareggio
pa = wb.create_sheet("Pareggio")
pa["A1"] = "Quanti Pass servono per coprire i costi fissi (scenario scelto)"
pa["A1"].font = TITOLO
pa["A2"] = ("Margine per Pass = prezzo medio meno costi variabili per Pass, al netto di contributi e imposta "
            "sulla parte variabile. In COM si aggiunge il contributo minimo fisso.")
pa["A2"].font = Font(name=FONT, italic=True, color="666666")
for j, t in enumerate(["Voce", "", "Gestione Separata", "Commercianti"], start=1):
    c = pa.cell(row=4, column=j, value=t)
    c.font = GRASSETTO
    c.fill = GRIGIO
P = [
    (5, "Prezzo medio incassato per Pass", EURO2, [f"={Asel('prezzo_medio')}"] * 2),
    (6, "Costi variabili per Pass (quota sui ricavi del 2027)", EURO2,
     [f"=IF({Asel('ricavi')}=0,0,C5*{Asel('c_var')}/{Asel('ricavi')})", "=C6"]),
    (7, "Contributi variabili per Pass", EURO2,
     [f"=C5*(1-{Isel('q_recessi')})*{Isel('coeff')}*{Isel('aliq_gs')}", "=0"]),
    (8, "Imposta per Pass", EURO2,
     [f"=(C5*(1-{Isel('q_recessi')})*{Isel('coeff')}-C7)*{Isel('imposta')}",
      f"=D5*(1-{Isel('q_recessi')})*{Isel('coeff')}*{Isel('imposta')}"]),
    (9, "Margine netto per Pass", EURO2, ["=C5-C6-C7-C8", "=D5-D6-D7-D8"]),
    (10, "Costi fissi annui", EURO, [f"={Asel('c_fissi')}", f"={Asel('c_fissi')}+{Isel('c_impresa')}"]),
    (11, "Contributo INPS fisso annuo", EURO, ["=0", f"={Isel('min_com')}*(1-{Isel('rid_forf')})"]),
    (12, "Pass l'anno per andare in pari", NUM, ["=IF(C9<=0,0,(C10+C11)/C9)", "=IF(D9<=0,0,(D10+D11)/D9)"]),
    (13, "Pass l'anno per l'obiettivo personale", NUM,
     [f"=IF(C9<=0,0,(C10+C11+{Isel('obiettivo')})/C9)", f"=IF(D9<=0,0,(D10+D11+{Isel('obiettivo')})/D9)"]),
    (14, "Confronto: mercato del recupero l'anno (persone)", NUM, [f"={Asel('mercato')}"] * 2),
]
for rr, et, fmt, formule in P:
    forte = rr in (9, 12, 13)
    pa.cell(row=rr, column=1, value=et).font = GRASSETTO if forte else NERO
    for j, f in enumerate(formule):
        c = pa.cell(row=rr, column=3 + j, value=f)
        c.number_format = fmt
        c.font = GRASSETTO if forte else NERO
pa.column_dimensions["A"].width = 55
pa.column_dimensions["B"].width = 2
pa.column_dimensions["C"].width = 18
pa.column_dimensions["D"].width = 18

for foglio in wb.worksheets:
    for riga in foglio.iter_rows():
        for cella in riga:
            if cella.font and cella.font.name != FONT:
                cella.font = Font(name=FONT, bold=cella.font.bold, italic=cella.font.italic,
                                  color=cella.font.color, size=cella.font.size)

wb.calculation.fullCalcOnLoad = True
wb.save(USCITA)
print(USCITA)
